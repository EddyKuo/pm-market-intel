#!/usr/bin/env python3
"""把市調報告 Markdown 轉成 Word (.docx)。

支援：標題 (#~####)、段落、**粗體**、*斜體*、`code`、[連結](url)（建立為可點擊超連結）、
無序/有序清單（含一層縮排）、核取方塊、GFM 表格（支援 \\| 跳脫、表頭跨頁重複、
欄數 ≥ --landscape-cols 時自動放進橫向頁面）、程式碼區塊、引用區塊、水平線。

用法：python3 build_docx.py report.md [-o report.docx] [--font "Microsoft JhengHei"] [--landscape-cols 7]
依賴：pip install python-docx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

INLINE = re.compile(
    r"(\[(?P<ltext>[^\]]+)\]\((?P<lurl>[^)\s]+)\)"   # [text](url)
    r"|\*\*(?P<bold>[^*]+)\*\*"                       # **bold**
    r"|`(?P<code>[^`]+)`"                             # `code`
    r"|(?<![*\w])\*(?P<ital>[^*\s][^*]*)\*(?!\*)"     # *italic*
    r"|(?P<url>https?://[^\s)<>|]+))"                 # bare url
)
LINK_COLOR = RGBColor(0x1F, 0x5F, 0xBF)


def set_font(run, font):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def add_hyperlink(par, text, url, font, size=None):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(attr), font)
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FBF")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_inline(par, text, font, size=None, bold=False):
    text = text.replace("\\|", "|").replace("\\*", "∗")
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _plain(par, text[pos:m.start()], font, size, bold)
        if m.group("ltext") is not None:
            add_hyperlink(par, m.group("ltext"), m.group("lurl"), font, size)
        elif m.group("url") is not None:
            add_hyperlink(par, m.group("url"), m.group("url"), font, size)
        elif m.group("bold") is not None:
            _plain(par, m.group("bold"), font, size, True)
        elif m.group("code") is not None:
            run = par.add_run(m.group("code"))
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x80, 0x30, 0x30)
            if size:
                run.font.size = Pt(size)
        elif m.group("ital") is not None:
            run = _plain(par, m.group("ital"), font, size, bold)
            run.italic = True
        pos = m.end()
    if pos < len(text):
        _plain(par, text[pos:], font, size, bold)


def _plain(par, text, font, size, bold):
    run = par.add_run(text.replace("∗", "*"))
    set_font(run, font)
    if size:
        run.font.size = Pt(size)
    if bold:
        run.bold = True
    return run


def parse_row(line):
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|") and not body.endswith("\\|"):
        body = body[:-1]
    cells = re.split(r"(?<!\\)\|", body)
    return [c.strip() for c in cells]


def is_separator(line):
    return bool(re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?", line.strip()))


def set_orientation(section, landscape: bool):
    w, h = section.page_width, section.page_height
    if landscape and w < h or not landscape and w > h:
        section.page_width, section.page_height = h, w
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_table(doc, rows, font, landscape_cols):
    ncols = max(len(r) for r in rows)
    wide = ncols >= landscape_cols
    if wide:
        set_orientation(doc.add_section(WD_SECTION.NEW_PAGE), True)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 8.5 if wide else 9.5
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            add_inline(cell.paragraphs[0], row[j] if j < len(row) else "", font, size, bold=(i == 0))
    repeat_header(table.rows[0])
    if wide:
        set_orientation(doc.add_section(WD_SECTION.NEW_PAGE), False)
    else:
        doc.add_paragraph()


def convert(md_path: Path, out_path: Path, font: str, landscape_cols: int = 7):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    for lvl in range(1, 5):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = font
        st.element.rPr.rFonts.set(qn("w:eastAsia"), font)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            par = doc.add_paragraph()
            run = par.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [parse_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_row(lines[i]))
                i += 1
            add_table(doc, rows, font, landscape_cols)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            heading = doc.add_heading(level=0 if (level == 1 and first_heading) else level)
            first_heading = False
            add_inline(heading, m.group(2), font)
            i += 1
            continue

        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            doc.add_paragraph("─" * 30)
            i += 1
            continue

        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            par = doc.add_paragraph(style="Intense Quote")
            add_inline(par, " ".join(b for b in buf if b), font)
            continue

        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if m:
            indent = len(m.group(1).replace("\t", "    "))
            numbered = m.group(2)[0].isdigit()
            style = "List Number" if numbered else "List Bullet"
            if indent >= 2:
                style += " 2"
            text = re.sub(r"^\[( |x|X)\]\s*", lambda x: "☑ " if x.group(1) in "xX" else "☐ ", m.group(3))
            add_inline(doc.add_paragraph(style=style), text, font)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^\s*(#{1,4}\s|[-*]\s|\d+\.\s|\||```|>)", lines[i]
        ):
            buf.append(lines[i].strip())
            i += 1
        add_inline(doc.add_paragraph(), " ".join(buf), font)

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("markdown")
    ap.add_argument("-o", "--output")
    ap.add_argument("--font", default="Microsoft JhengHei")
    ap.add_argument("--landscape-cols", type=int, default=7, help="欄數達此值的表格放進橫向頁面")
    args = ap.parse_args()
    md = Path(args.markdown)
    out = Path(args.output) if args.output else md.with_suffix(".docx")
    print(convert(md, out, args.font, args.landscape_cols))


if __name__ == "__main__":
    main()
