"""pm-market-intel 腳本測試。執行：python3 -m unittest discover -s tests -v"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "plugins/pm-market-intel/skills/market-intel/scripts"
sys.path.insert(0, str(SCRIPTS))
import mi_common as mc  # noqa: E402


def run(script, *args):
    return subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)],
                          capture_output=True, text=True)


def dump(path: Path, data):
    path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def player(pid="P-0001", **kw):
    base = {
        "id": pid, "name": f"Co {pid}", "roles": ["brand"], "primary_lens": "成本", "regions": ["TW"],
        "segments": ["S-01"], "depth": "profile", "evidence": ["E-0001"],
        "first_seen": "2026-10-01", "last_verified": "2026-10-01", "status": "active",
    }
    base.update(kw)
    return base


def ev(eid="E-0001", **kw):
    base = {
        "id": eid, "claim": f"claim {eid}", "claim_type": "vendor_claim", "source_type": "official",
        "source": "site", "url": f"https://example.com/{eid}", "origin": "Vendor", "accessed": "2026-10-01",
        "confidence": "A", "related": ["P-0001"],
    }
    base.update(kw)
    return base


def opp(oid="O-0001", **kw):
    base = {
        "id": oid, "hypothesis": "h", "segment": "S-01", "stakeholder": "user", "lens": "功能",
        "gap_type": "supply", "evidence_for": ["E-0001"], "evidence_against": [], "notes": "searched X",
        "fit": {"score": 3, "rationale": "r"}, "scores": {"problem_importance": 4, "unmet": 3},
        "evidence_strength": "weak", "kill_criteria": ["k"],
        "validation_next": [{"method": "interview", "question": "q"}], "status": "open",
    }
    base.update(kw)
    return base


class ProjectCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.proj = self.tmp / "proj"
        self.proj.mkdir()
        (self.proj / "10-value-map.md").write_text("# 價值地圖\n\n## S-01 中型工廠\n", encoding="utf-8")
        dump(self.proj / "players.yaml", [player()])
        dump(self.proj / "evidence.yaml", [ev()])
        dump(self.proj / "opportunities.yaml", [opp()])

    def tearDown(self):
        shutil.rmtree(self.tmp)


class TestDiff(ProjectCase):
    def snapshot_and_change(self, mutate):
        old = self.tmp / "old"
        shutil.copytree(self.proj, old)
        players = mc.load_list(self.proj / "players.yaml")
        mutate(players)
        dump(self.proj / "players.yaml", players)
        return run("diff_landscape.py", old, self.proj)

    def test_reactivation_is_reported(self):
        dump(self.proj / "players.yaml", [player(status="inactive", status_note="停售")])

        def mutate(ps):
            ps[0]["status"] = "active"
            ps[0]["status_note"] = "重新上市"
        r = self.snapshot_and_change(mutate)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("重新啟用 1", r.stdout)
        self.assertIn("inactive → active", r.stdout)
        self.assertIn("20-landscape.md", r.stdout)

    def test_exit_is_reported(self):
        def mutate(ps):
            ps[0]["status"] = "inactive"
        r = self.snapshot_and_change(mutate)
        self.assertIn("退出 1", r.stdout)
        self.assertIn("active → inactive", r.stdout)

    def test_duplicate_id_fails_loudly(self):
        old = self.tmp / "old"
        shutil.copytree(self.proj, old)
        dump(self.proj / "players.yaml", [player(), player(name="Other")])
        r = run("diff_landscape.py", old, self.proj)
        self.assertEqual(r.returncode, 2)
        self.assertIn("重複 ID", r.stderr)

    def test_evidence_change_flags_opportunity(self):
        old = self.tmp / "old"
        shutil.copytree(self.proj, old)
        dump(self.proj / "evidence.yaml", [ev(confidence="C")])
        r = run("diff_landscape.py", old, self.proj)
        self.assertIn("opportunities.yaml", r.stdout)
        self.assertIn("O-0001", r.stdout)

    def test_player_exit_flags_dependent_opportunity(self):
        def mutate(ps):
            ps[0]["status"] = "inactive"
            ps[0]["status_note"] = "停止服務"
        r = self.snapshot_and_change(mutate)
        self.assertIn("證據所涉 player 狀態或資料變動：P-0001", r.stdout)

    def test_pipe_in_value_is_escaped(self):
        def mutate(ps):
            ps[0]["business_model"] = "A | B"
        r = self.snapshot_and_change(mutate)
        self.assertIn("A \\| B", r.stdout)


class TestValidate(ProjectCase):
    def test_valid_project_passes(self):
        r = run("validate.py", self.proj)
        self.assertEqual(r.returncode, 0, r.stdout)

    def test_duplicate_id(self):
        dump(self.proj / "players.yaml", [player(), player()])
        r = run("validate.py", self.proj)
        self.assertEqual(r.returncode, 1)
        self.assertIn("重複 ID", r.stdout)

    def test_dangling_reference_in_yaml(self):
        dump(self.proj / "players.yaml", [player(evidence=["E-0099"])])
        r = run("validate.py", self.proj)
        self.assertIn("E-0099", r.stdout)
        self.assertEqual(r.returncode, 1)

    def test_dangling_reference_in_markdown(self):
        (self.proj / "30-comparison.md").write_text("依據 E-0042 與 P-0001", encoding="utf-8")
        r = run("validate.py", self.proj)
        self.assertIn("E-0042", r.stdout)
        self.assertNotIn("`P-0001`", r.stdout)

    def test_bad_enum_and_single_role_string(self):
        dump(self.proj / "players.yaml", [player(roles="brand", primary_lens="品牌")])
        r = run("validate.py", self.proj)
        self.assertIn("`roles` 必須是 list", r.stdout)
        self.assertIn("primary_lens", r.stdout)

    def test_overclaimed_strength(self):
        dump(self.proj / "opportunities.yaml", [opp(evidence_strength="strong")])
        r = run("validate.py", self.proj)
        self.assertIn("依規則只有 weak", r.stdout)

    def test_missing_evidence_against(self):
        o = opp()
        del o["evidence_against"]
        dump(self.proj / "opportunities.yaml", [o])
        r = run("validate.py", self.proj)
        self.assertIn("evidence_against", r.stdout)

    def test_teardown_file_must_exist(self):
        dump(self.proj / "players.yaml", [player(depth="teardown", teardown="teardowns/P-0001.md")])
        r = run("validate.py", self.proj)
        self.assertIn("不存在", r.stdout)
        (self.proj / "teardowns").mkdir()
        (self.proj / "teardowns/P-0001.md").write_text("# teardown P-0001", encoding="utf-8")
        self.assertEqual(run("validate.py", self.proj).returncode, 0)


class TestStrength(unittest.TestCase):
    def test_rules(self):
        evidence = {
            "E-1": {"id": "E-1", "claim_type": "vendor_claim", "confidence": "A", "origin": "Vendor"},
            "E-2": {"id": "E-2", "claim_type": "vendor_claim", "confidence": "B", "origin": "Vendor"},
            "E-3": {"id": "E-3", "claim_type": "measured", "confidence": "B", "origin": "Lab"},
            "E-4": {"id": "E-4", "claim_type": "user_voice", "confidence": "A", "origin": "Forum"},
        }
        self.assertEqual(mc.compute_strength(["E-1", "E-2"], [], evidence), "weak")   # 同源轉載
        self.assertEqual(mc.compute_strength(["E-3"], [], evidence), "moderate")
        self.assertEqual(mc.compute_strength(["E-1", "E-3"], [], evidence), "strong")
        self.assertEqual(mc.compute_strength(["E-1", "E-3"], ["E-4"], evidence), "moderate")  # 有強反證


class TestMerge(ProjectCase):
    def test_merge_assigns_ids_and_dedupes(self):
        dump(self.proj / "players.yaml", [player(website="https://www.alpha.com", name="Alpha Inc.")])
        b1 = self.tmp / "b1.yaml"
        b2 = self.tmp / "b2.yaml"
        dump(b1, {"scout": "brand", "players": [
            player("tmp-p1", name="Alpha", website="alpha.com", roles=["integrator"], evidence=["tmp-e1"]),
            player("tmp-p2", name="Beta Co., Ltd.", evidence=["tmp-e2"]),
        ], "evidence": [ev("tmp-e1", related=["tmp-p1"]), ev("tmp-e2", related=["tmp-p2"])]})
        dump(b2, {"scout": "integrator", "players": [
            player("tmp-p1", name="Beta", evidence=["tmp-e1"]),
        ], "evidence": [ev("tmp-e1", url="https://example.com/tmp-e2", claim="claim tmp-e2", related=["tmp-p1"])]})
        r = run("merge_scout.py", self.proj, b1, b2)
        self.assertEqual(r.returncode, 0, r.stderr)
        players = mc.index_by_id(mc.load_list(self.proj / "players.yaml"), "p")
        evidence = mc.index_by_id(mc.load_list(self.proj / "evidence.yaml"), "e")
        self.assertEqual(sorted(players), ["P-0001", "P-0002"])
        self.assertEqual(players["P-0001"]["roles"], ["brand", "integrator"])
        self.assertEqual(sorted(evidence), ["E-0001", "E-0002", "E-0003"])
        self.assertEqual(players["P-0002"]["evidence"], ["E-0003"])
        self.assertEqual(evidence["E-0003"]["related"], ["P-0002"])
        self.assertIn("evidence_dup=1", r.stdout)


class TestSnapshot(ProjectCase):
    def test_unique_runs_and_previous(self):
        r1 = run("snapshot.py", self.proj, "--label", "first")
        r2 = run("snapshot.py", self.proj, "--label", "first")
        self.assertEqual(r1.returncode, 0, r1.stderr)
        id1 = r1.stdout.split("run_id=")[1].split()[0]
        id2 = r2.stdout.split("run_id=")[1].split()[0]
        self.assertNotEqual(id1, id2)
        self.assertIn(id1, r2.stdout.split("previous=")[1])
        self.assertTrue((self.proj / "runs" / id2 / "players.yaml").exists())
        latest = run("snapshot.py", self.proj, "--latest")
        self.assertIn(id2, latest.stdout)


class TestDocx(unittest.TestCase):
    def test_links_pipes_landscape(self):
        import docx
        tmp = Path(tempfile.mkdtemp())
        md = tmp / "r.md"
        md.write_text(textwrap.dedent("""\
            # 報告

            見 [官網](https://example.com/a) 與 https://example.com/b 。

            | A | B |
            |---|---|
            | x \\| y | **粗** |

            | 1 | 2 | 3 | 4 | 5 | 6 | 7 |
            |---|---|---|---|---|---|---|
            | a | b | c | d | e | f | g |
            """), encoding="utf-8")
        r = run("build_docx.py", md)
        self.assertEqual(r.returncode, 0, r.stderr)
        d = docx.Document(tmp / "r.docx")
        rels = [rel.target_ref for rel in d.part.rels.values() if "hyperlink" in rel.reltype]
        self.assertIn("https://example.com/a", rels)
        self.assertIn("https://example.com/b", rels)
        t0 = d.tables[0]
        self.assertEqual(len(t0.columns), 2)
        self.assertEqual(t0.cell(1, 0).text, "x | y")
        self.assertIn("w:tblHeader", t0.rows[0]._tr.xml)
        orientations = [s.orientation for s in d.sections]
        self.assertIn(1, orientations)        # 有橫向頁
        self.assertEqual(orientations[-1], 0)  # 最後回到直向
        shutil.rmtree(tmp)


if __name__ == "__main__":
    unittest.main()
