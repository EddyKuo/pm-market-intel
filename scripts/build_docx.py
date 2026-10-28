#!/usr/bin/env python3
"""把市調報告 Markdown 轉成 Word (.docx)。

支援：標題 (#~####)、段落、**粗體**、`code`、無序/有序清單（含一層縮排）、
表格（GFM pipe table）、程式碼區塊、水平線。

用法：python build_docx.py report.md [-o report.docx] [--font "Microsoft JhengHei"]
依賴：pip install python-docx
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def set_font(run, font):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def add_inline(par, text, font):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x80, 0x30, 0x30)
            continue
        else:
            run = par.add_run(part)
        set_font(run, font)


def parse_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def is_separator(line):
    return bool(re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?", line.strip()))


def add_table(doc, rows, font):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            add_inline(par, row[j] if j < len(row) else "", font)
            for run in par.runs:
                run.font.size = Pt(9.5)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def convert(md_path: Path, out_path: Path, font: str):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    for lvl in range(1, 5):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = font
        st.element.rPr.rFonts.set(qn("w:eastAsia"), font)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            par = doc.add_paragraph()
            run = par.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [parse_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_row(lines[i]))
                i += 1
            add_table(doc, rows, font)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            if level == 1 and not doc.paragraphs:
                heading = doc.add_heading(level=0)
            else:
                heading = doc.add_heading(level=level)
            add_inline(heading, m.group(2), font)
            i += 1
            continue

        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            doc.add_paragraph("─" * 30)
            i += 1
            continue

        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if m:
            indent = len(m.group(1).replace("\t", "    "))
            numbered = m.group(2)[0].isdigit()
            style = "List Number" if numbered else "List Bullet"
            if indent >= 2:
                style += " 2"
            text = re.sub(r"^\[( |x)\]\s*", lambda x: "☑ " if x.group(1) == "x" else "☐ ", m.group(3))
            add_inline(doc.add_paragraph(style=style), text, font)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^\s*(#{1,4}\s|[-*]\s|\d+\.\s|\||```)", lines[i]
        ):
            buf.append(lines[i].strip())
            i += 1
        add_inline(doc.add_paragraph(), " ".join(buf), font)

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("markdown")
    ap.add_argument("-o", "--output")
    ap.add_argument("--font", default="Microsoft JhengHei")
    args = ap.parse_args()
    md = Path(args.markdown)
    out = Path(args.output) if args.output else md.with_suffix(".docx")
    print(convert(md, out, args.font))


if __name__ == "__main__":
    main()
